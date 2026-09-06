"""
firebase_db.py — Firestore Database Layer
Drop-in replacement for models.py using Firebase Firestore.
All function signatures and return formats are identical to models.py.
"""
import csv
import os
from datetime import datetime, date, timedelta
from firebase_client import get_db
from config import EXPORT_DIR


# ── Helpers ────────────────────────────────────────────────────────────────────
def _col(name):
    return get_db().collection(name)


def _doc_to_dict(doc):
    if doc.exists:
        d = doc.to_dict()
        d["_id"] = doc.id
        return d
    return None


# ══════════════════════════════════════════════════════════════
# INIT
# ══════════════════════════════════════════════════════════════
def init_db():
    """Verify Firestore connection (collections auto-create on first write)."""
    try:
        get_db()
        print("[Firebase] Firestore ready")
    except Exception as e:
        raise ConnectionError(f"[Firebase] Connection failed: {e}")


# ══════════════════════════════════════════════════════════════
# STUDENTS
def normalize_course(course: str, year: str, department: str = "") -> str:
    c = str(course or "").strip()
    y = str(year or "").strip().lower()
    d = str(department or "").strip()
    if y.startswith("pg"):
        if not c or c == "B.Tech":
            return "M.Tech"
        return c
    if not c:
        if d in ["Finance", "Marketing"]:
            return "MBA"
        return "B.Tech"
    return c


def add_student(roll_no, name, email="", department="", year="", academic_year="", course=""):
    course = normalize_course(course, year, department)
    ref = _col("students").document(roll_no)
    if ref.get().exists:
        return False
    ref.set({
        "roll_no": roll_no, "name": name, "email": email,
        "course": course, "department": department, "year": year, "academic_year": academic_year,
        "enrolled_at": datetime.now().isoformat(),
    })
    return True


def update_student(roll_no, name, email="", department="", year="", academic_year="", course=""):
    course = normalize_course(course, year, department)
    _col("students").document(roll_no).update({
        "name": name, "email": email,
        "course": course, "department": department, "year": year, "academic_year": academic_year,
    })


def get_all_students():
    docs = _col("students").order_by("roll_no").stream()
    return [d.to_dict() for d in docs]


def get_student(roll_no):
    doc = _col("students").document(roll_no).get()
    return doc.to_dict() if doc.exists else None


def delete_student_record(roll_no):
    # Delete all attendance for this student
    for a in _col("attendance").where("roll_no", "==", roll_no).stream():
        a.reference.delete()
    _col("students").document(roll_no).delete()


# ══════════════════════════════════════════════════════════════
# SUBJECTS
# ══════════════════════════════════════════════════════════════
def add_subject(code, name, department="", credits=3, classes=None):
    # Check for duplicate code
    existing = _col("subjects").where("code", "==", code.upper()).limit(1).stream()
    if any(True for _ in existing):
        return False
    _col("subjects").add({
        "code": code.upper(), "name": name,
        "department": department, "credits": credits,
        "classes": classes or []
    })
    return True


def update_subject_mappings(subject_id, classes):
    # Try ID first
    doc = _col("subjects").document(subject_id).get()
    if doc.exists:
        doc.reference.update({"classes": classes or []})
        return True
    
    # Try by code
    existing = _col("subjects").where("code", "==", subject_id).limit(1).stream()
    for d in existing:
        d.reference.update({"classes": classes or []})
        return True
    return False


def get_all_subjects():
    docs = _col("subjects").order_by("code").stream()
    result = []
    for d in docs:
        row = d.to_dict()
        row["id"] = d.id
        result.append(row)
    return result


def get_subject(subject_id):
    doc = _col("subjects").document(subject_id).get()
    if not doc.exists:
        return None
    row = doc.to_dict()
    row["id"] = doc.id
    return row


def delete_subject(subject_id):
    # Delete sessions and their attendance
    for ses in _col("sessions").where("subject_id", "==", subject_id).stream():
        for a in _col("attendance").where("session_id", "==", ses.id).stream():
            a.reference.delete()
        ses.reference.delete()
    _col("subjects").document(subject_id).delete()


# ══════════════════════════════════════════════════════════════
# SESSIONS
# ══════════════════════════════════════════════════════════════
def create_session(subject_id, faculty="", notes="", department="", year="", academic_year=""):
    today = date.today().isoformat()
    now   = datetime.now().strftime("%H:%M:%S")
    _, ref = _col("sessions").add({
        "subject_id": subject_id,
        "date": today, "start_time": now,
        "faculty": faculty, "notes": notes,
        "department": department,
        "year": year,
        "academic_year": academic_year
    })
    return ref.id   # string ID


def get_all_sessions():
    sessions = _col("sessions").order_by("date", direction="DESCENDING").stream()
    result = []
    for s in sessions:
        row = s.to_dict()
        row["id"] = s.id
        # Fetch subject info
        sub = _col("subjects").document(row["subject_id"]).get()
        if sub.exists:
            sd = sub.to_dict()
            row["code"] = sd.get("code", "?")
            row["subject_name"] = sd.get("name", "?")
        result.append(row)
    return result


def get_sessions_for_subject(subject_id):
    sessions = (_col("sessions")
                .where("subject_id", "==", subject_id)
                .stream())
    result = []
    sub = _col("subjects").document(subject_id).get()
    sub_data = sub.to_dict() if sub.exists else {}
    for s in sessions:
        row = s.to_dict()
        row["id"] = s.id
        row["code"] = sub_data.get("code", "?")
        row["subject_name"] = sub_data.get("name", "?")
        result.append(row)
    result.sort(key=lambda x: x.get("date", ""), reverse=True)
    return result


def count_sessions(subject_id=None):
    if subject_id:
        docs = _col("sessions").where("subject_id", "==", subject_id).stream()
    else:
        docs = _col("sessions").stream()
    return sum(1 for _ in docs)


# ══════════════════════════════════════════════════════════════
# ATTENDANCE
# ══════════════════════════════════════════════════════════════
def mark_attendance(roll_no, session_id, confidence=None):
    today = date.today().isoformat()
    now   = datetime.now().strftime("%H:%M:%S")
    
    # Fetch subject_id from session to enforce uniqueness per subject per day
    session_doc = _col("sessions").document(session_id).get()
    subject_id  = session_doc.to_dict().get("subject_id") if session_doc.exists else "unknown"
    
    # Composite doc ID enforces uniqueness (roll_no + subject_id + date)
    doc_id = f"{roll_no}__{subject_id}__{today}"
    ref    = _col("attendance").document(doc_id)
    if ref.get().exists:
        return "duplicate"
    try:
        ref.set({
            "roll_no": roll_no, "session_id": session_id,
            "date": today, "time": now, "confidence": confidence,
        })
        return "marked"
    except Exception as e:
        print(f"[Firebase] mark_attendance error: {e}")
        return "error"


def _enrich_attendance(rows):
    """Add student name/dept and subject code/name to raw attendance rows."""
    db = get_db()
    # Batch-fetch unique students using db.get_all
    roll_nos = list({r["roll_no"] for r in rows})
    students = {}
    if roll_nos:
        refs = [db.collection("students").document(rn) for rn in roll_nos]
        docs = db.get_all(refs)
        for doc in docs:
            if doc.exists:
                students[doc.id] = doc.to_dict()

    # Batch-fetch unique sessions → subjects
    session_ids = list({r["session_id"] for r in rows})
    sessions, subjects = {}, {}
    if session_ids:
        s_refs = [db.collection("sessions").document(sid) for sid in session_ids]
        s_docs = db.get_all(s_refs)
        
        subj_ids_to_fetch = set()
        for doc in s_docs:
            if doc.exists:
                sdata = doc.to_dict()
                sessions[doc.id] = sdata
                if sdata.get("subject_id"):
                    subj_ids_to_fetch.add(sdata.get("subject_id"))
        
        if subj_ids_to_fetch:
            sub_refs = [db.collection("subjects").document(sid) for sid in subj_ids_to_fetch]
            sub_docs = db.get_all(sub_refs)
            for doc in sub_docs:
                if doc.exists:
                    subjects[doc.id] = doc.to_dict()

    result = []
    for r in rows:
        rn  = r["roll_no"]
        sid = r["session_id"]
        st  = students.get(rn, {})
        ses = sessions.get(sid, {})
        sub = subjects.get(ses.get("subject_id", ""), {})
        result.append({
            "roll_no":      rn,
            "name":         st.get("name", rn),
            "department":   st.get("department", ""),
            "year":         st.get("year", ""),
            "academic_year": st.get("academic_year", ""),
            "subject_code": sub.get("code", "?"),
            "subject_name": sub.get("name", "?"),
            "subject_id":   ses.get("subject_id"), # Added for filtering
            "session_id":   sid,                   # Added for completeness
            "date":         r.get("date", ""),
            "time":         r.get("time", ""),
            "confidence":   r.get("confidence"),
        })
    return result


def get_session_attendance(session_id):
    docs = _col("attendance").where("session_id", "==", session_id).stream()
    rows = [d.to_dict() for d in docs]
    rows.sort(key=lambda x: x.get("time", ""))
    return _enrich_attendance(rows)


def get_today_attendance(subject_id=None):
    today = date.today().isoformat()
    q = _col("attendance").where("date", "==", today)
    rows = [d.to_dict() for d in q.stream()]
    enriched = _enrich_attendance(rows)
    if subject_id:
        enriched = [r for r in enriched
                    if r.get("subject_id") == subject_id]
    enriched.sort(key=lambda x: x.get("time", ""))
    return enriched


def _get_subject_id_for_session(row):
    """Helper for filtering by subject_id via session."""
    sid = row.get("session_id")
    if not sid:
        return None
    doc = _col("sessions").document(sid).get()
    return doc.to_dict().get("subject_id") if doc.exists else None


def get_attendance_by_date(query_date, subject_id=None):
    rows = [d.to_dict() for d in
            _col("attendance").where("date", "==", query_date).stream()]
    enriched = _enrich_attendance(rows)
    if subject_id:
        enriched = [r for r in enriched
                    if r.get("subject_id") == subject_id]
    enriched.sort(key=lambda x: x.get("time", ""))
    return enriched


def get_student_history(roll_no):
    docs = (_col("attendance")
            .where("roll_no", "==", roll_no)
            .stream())
    rows = [d.to_dict() for d in docs]
    enriched = _enrich_attendance(rows)
    enriched.sort(key=lambda x: (x.get("date", ""), x.get("time", "")), reverse=True)
    return enriched


def get_attendance_summary(subject_id=None):
    students = get_all_students()
    total = count_sessions(subject_id)

    # Fetch all relevant attendance
    if subject_id:
        session_ids = {s["id"] for s in get_sessions_for_subject(subject_id)}
        att_docs = _col("attendance").stream()
        att_by_roll = {}
        for d in att_docs:
            row = d.to_dict()
            if row.get("session_id") in session_ids:
                att_by_roll[row["roll_no"]] = att_by_roll.get(row["roll_no"], 0) + 1
    else:
        att_docs = _col("attendance").stream()
        att_by_roll = {}
        for d in att_docs:
            rn = d.to_dict()["roll_no"]
            att_by_roll[rn] = att_by_roll.get(rn, 0) + 1

    result = []
    for s in students:
        present = att_by_roll.get(s["roll_no"], 0)
        pct = round((present / total * 100) if total > 0 else 0, 1)
        result.append({
            "roll_no":       s["roll_no"],
            "name":          s["name"],
            "department":    s.get("department", ""),
            "year":          s.get("year", ""),
            "days_present":  present,
            "total_sessions": total,
            "pct":           pct,
        })
    result.sort(key=lambda x: x["days_present"], reverse=True)
    return result


def get_daily_trend(subject_id):
    """Return list of {'date': 'YYYY-MM-DD', 'count': N} for a subject."""
    sessions = get_sessions_for_subject(subject_id)
    session_ids = {s["id"] for s in sessions}
    
    docs = _col("attendance").stream()
    counts_by_date = {}
    
    for d in docs:
        row = d.to_dict()
        if row.get("session_id") in session_ids:
            dt = row.get("date")
            counts_by_date[dt] = counts_by_date.get(dt, 0) + 1
            
    # Sort by date
    dates = sorted(counts_by_date.keys())
    return [{"date": d, "count": counts_by_date[d]} for d in dates]


# ══════════════════════════════════════════════════════════════
# EXPORT
# ══════════════════════════════════════════════════════════════
def _get_export_rows(from_date=None, to_date=None, subject_id=None):
    docs = _col("attendance").stream()
    rows = [d.to_dict() for d in docs]

    if from_date:
        rows = [r for r in rows if r.get("date", "") >= from_date]
    if to_date:
        rows = [r for r in rows if r.get("date", "") <= to_date]

    enriched = _enrich_attendance(rows)

    if subject_id:
        all_sids = {s["id"] for s in get_sessions_for_subject(subject_id)}
        enriched = [r for r in enriched
                    if any(d.to_dict().get("session_id") in all_sids
                           for d in _col("attendance")
                           .where("roll_no", "==", r["roll_no"])
                           .where("date", "==", r["date"]).stream())]

    enriched.sort(key=lambda x: (x.get("date", ""), x.get("time", "")), reverse=True)
    return enriched


def export_to_csv(output_path=None, from_date=None, to_date=None, subject_id=None):
    os.makedirs(EXPORT_DIR, exist_ok=True)
    if not output_path:
        output_path = os.path.join(EXPORT_DIR, f"attendance_{date.today().isoformat()}.csv")
    rows = _get_export_rows(from_date, to_date, subject_id)
    with open(output_path, "w", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["Roll No","Name","Department","Year",
                         "Subject Code","Subject","Date","Time","Confidence"])
        for r in rows:
            writer.writerow([r["roll_no"], r["name"], r["department"], r["year"],
                             r["subject_code"], r["subject_name"],
                             r["date"], r["time"], r.get("confidence","")])
    return output_path


def export_to_excel(output_path=None, from_date=None, to_date=None, subject_id=None):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    os.makedirs(EXPORT_DIR, exist_ok=True)
    if not output_path:
        output_path = os.path.join(EXPORT_DIR, f"attendance_{date.today().isoformat()}.xlsx")

    rows = _get_export_rows(from_date, to_date, subject_id)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance"

    headers = ["Roll No","Name","Department","Year",
               "Subject Code","Subject","Date","Time","Confidence"]
    hfill = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
    hfont = Font(color="FFFFFF", bold=True)
    thin  = Side(style="thin", color="D1D5DB")
    bdr   = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.fill = hfill; c.font = hfont
        c.alignment = Alignment(horizontal="center"); c.border = bdr

    for ri, row in enumerate(rows, 2):
        vals = [row["roll_no"], row["name"], row["department"], row["year"],
                row["subject_code"], row["subject_name"],
                row["date"], row["time"], row.get("confidence","")]
        fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid") \
               if ri % 2 == 0 else None
        for ci, val in enumerate(vals, 1):
            c = ws.cell(row=ri, column=ci, value=val)
            c.border = bdr
            if fill: c.fill = fill

    for i, w in enumerate([12,22,14,14,14,24,12,10,12], 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    wb.save(output_path)
    return output_path


def delete_attendance(roll_no, session_id):
    docs = _col("attendance").where("roll_no", "==", roll_no).where("session_id", "==", session_id).stream()
    deleted = False
    for doc in docs:
        doc.reference.delete()
        deleted = True
    return deleted


def update_attendance(old_roll_no, session_id, new_roll_no, time_str):
    # Verify student exists if roll_no is changing
    if old_roll_no != new_roll_no:
        student_doc = _col("students").document(new_roll_no).get()
        if not student_doc.exists:
            return False

    # Fetch the session to get subject_id
    session_doc = _col("sessions").document(session_id).get()
    if not session_doc.exists:
        return False
    subject_id = session_doc.to_dict().get("subject_id", "unknown")

    # Find old attendance document
    docs = _col("attendance").where("roll_no", "==", old_roll_no).where("session_id", "==", session_id).limit(1).stream()
    old_doc = None
    for d in docs:
        old_doc = d
        break

    if not old_doc:
        return False

    old_data = old_doc.to_dict()
    date_val = old_data.get("date")
    confidence_val = old_data.get("confidence")

    # Check if new document ID already exists if roll_no is changing
    if old_roll_no != new_roll_no:
        new_doc_id = f"{new_roll_no}__{subject_id}__{date_val}"
        if _col("attendance").document(new_doc_id).get().exists:
            return "duplicate"

    # Save changes
    if old_roll_no == new_roll_no:
        old_doc.reference.update({
            "time": time_str
        })
    else:
        # Delete old document
        old_doc.reference.delete()

        # Create new document with composite ID
        new_doc_id = f"{new_roll_no}__{subject_id}__{date_val}"
        _col("attendance").document(new_doc_id).set({
            "roll_no": new_roll_no,
            "session_id": session_id,
            "date": date_val,
            "time": time_str,
            "confidence": confidence_val
        })
    return True


# ══════════════════════════════════════════════════════════════
# USERS
# ══════════════════════════════════════════════════════════════
def get_user(username):
    doc = _col("users").document(username).get()
    return doc.to_dict() if doc.exists else None

def add_user(username, password_hash, role, name, reference_id="", subjects=None, classes=None):
    ref = _col("users").document(username)
    if ref.get().exists:
        return False
    ref.set({
        "username": username,
        "password_hash": password_hash,
        "role": role,
        "name": name,
        "reference_id": reference_id,
        "subjects": subjects or [],
        "classes": classes or []
    })
    return True

def update_user(username, data):
    _col("users").document(username).update(data)
    return True

def delete_user(username):
    _col("users").document(username).delete()
    return True

def list_teachers():
    docs = _col("users").where("role", "==", "teacher").stream()
    return [d.to_dict() for d in docs]


# ══════════════════════════════════════════════════════════════
# SECTIONS  (Section-to-Subject Mapping)
# ══════════════════════════════════════════════════════════════
def _section_id(department, year, academic_year):
    """Generate a deterministic Firestore document ID for a section."""
    return f"{department}__{year}__{academic_year}".replace(" ", "_")


def add_section(department, year, academic_year, subjects=None):
    """Create a new section if it doesn't already exist. Returns True on success."""
    doc_id = _section_id(department, year, academic_year)
    ref = _col("sections").document(doc_id)
    if ref.get().exists:
        return False
    ref.set({
        "department": department,
        "year": year,
        "academic_year": academic_year,
        "subjects": subjects or [],
        "created_at": __import__("datetime").datetime.now().isoformat(),
    })
    return True


def get_all_sections():
    """Return all sections ordered by department."""
    docs = _col("sections").order_by("department").stream()
    result = []
    for d in docs:
        row = d.to_dict()
        row["id"] = d.id
        result.append(row)
    return result


def get_section(department, year, academic_year):
    """Fetch a single section by its composite key."""
    doc_id = _section_id(department, year, academic_year)
    doc = _col("sections").document(doc_id).get()
    if not doc.exists:
        return None
    row = doc.to_dict()
    row["id"] = doc.id
    return row


def get_section_by_id(section_id):
    """Fetch a single section by Firestore document ID."""
    doc = _col("sections").document(section_id).get()
    if not doc.exists:
        return None
    row = doc.to_dict()
    row["id"] = doc.id
    return row


def update_section_subjects(section_id, subjects):
    """Replace the subjects list for a section."""
    ref = _col("sections").document(section_id)
    if not ref.get().exists:
        return False
    ref.update({"subjects": subjects or []})
    return True


def delete_section(section_id):
    """Delete a section document."""
    _col("sections").document(section_id).delete()
    return True


def get_subjects_for_class(department, year, academic_year):
    """Return the list of subject codes mapped to the given class/section."""
    sec = get_section(department, year, academic_year)
    return sec["subjects"] if sec else []

def get_all_departments():
    docs = _col("departments").stream()
    return sorted([doc.id for doc in docs])

def add_department(code):
    ref = _col("departments").document(code.strip().upper())
    if ref.get().exists:
        return False
    ref.set({"code": code.strip().upper()})
    return True

def delete_department(code):
    _col("departments").document(code).delete()
    return True

def get_all_academic_years():
    docs = _col("academic_years").stream()
    return sorted([doc.id for doc in docs])

def add_academic_year(code):
    ref = _col("academic_years").document(code.strip())
    if ref.get().exists:
        return False
    ref.set({"code": code.strip()})
    return True

def delete_academic_year(code):
    _col("academic_years").document(code).delete()
    return True

def get_all_courses():
    docs = _col("courses").stream()
    codes = [doc.id for doc in docs]
    if not codes:
        return ["B.E. / B.Tech", "M.E. / M.Tech", "MCA", "MBA", "B.Sc", "M.Sc"]
    return sorted(codes)

def add_course(code):
    ref = _col("courses").document(code.strip())
    if ref.get().exists:
        return False
    ref.set({"code": code.strip()})
    return True

def delete_course(code):
    _col("courses").document(code).delete()
    return True

def get_all_schedules():
    docs = _col("schedules").stream()
    result = []
    # Fetch all subjects to resolve code and name
    subs = {str(s["id"]): s for s in get_all_subjects()}
    for doc in docs:
        d = doc.to_dict()
        d["id"] = doc.id
        sub_id = str(d.get("subject_id"))
        if sub_id in subs:
            d["subject_code"] = subs[sub_id]["code"]
            d["subject_name"] = subs[sub_id]["name"]
        else:
            d["subject_code"] = "UNKNOWN"
            d["subject_name"] = "Unknown Subject"
        result.append(d)
    
    # Sort helper
    days = ['Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday', 'Saturday', 'Sunday']
    result.sort(key=lambda x: (days.index(x.get("day_of_week", "Monday")) if x.get("day_of_week") in days else 0, x.get("start_time", "")))
    return result

def add_schedule(department, year, academic_year, subject_id, day_of_week, start_time, end_time, classroom="", specific_date=""):
    ref = _col("schedules").document()
    ref.set({
        "department": department,
        "year": year,
        "academic_year": academic_year,
        "subject_id": int(subject_id),
        "day_of_week": day_of_week,
        "start_time": start_time,
        "end_time": end_time,
        "classroom": classroom,
        "specific_date": specific_date
    })
    return True

def delete_schedule(schedule_id):
    _col("schedules").document(schedule_id).delete()
    return True
